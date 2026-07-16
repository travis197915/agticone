"""Excel ingestion: pull claim rows (id + optional billing columns) from `.xlsx`."""
from __future__ import annotations

import io
import logging
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Iterable

logger = logging.getLogger(__name__)

_CLAIM_ID_ALIASES = (
    "claim_id", "claimid", "subscriber_id", "subscriberid",
    "claim id", "subscriber id",
)

# Optional columns surfaced on the claim response for the UI.
_OPTIONAL_COLUMNS: dict[str, tuple[str, ...]] = {
    "paid_dt": ("paid_dt", "paid dt", "paid date"),
    "total_billed": ("total_billed", "total billed"),
    "total_paid": ("total_paid", "total paid"),
}

EXCEL_BILLING_FIELDS = tuple(_OPTIONAL_COLUMNS.keys())

# Kept separate from _OPTIONAL_COLUMNS/EXCEL_BILLING_FIELDS: this one feeds
# RuleExecutionRun.original_auditor (a dedicated model column resolved via
# execution_app.reviewer_lookup), not the generic claim-payload spread —
# mixing it into EXCEL_BILLING_FIELDS would let the raw sheet value clobber
# the resolved name in API responses.
_AUDITOR_NAME_ALIASES = ("auditorname", "auditor_name", "auditor name")


class XlsxParseError(ValueError):
    pass


def _norm_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "_")


def _cell_value(value: Any) -> Any:
    """Normalize Excel cell values for JSON / API payloads."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    text = str(value).strip()
    return text if text else None


def _find_column(norm_headers: list[str], aliases: tuple[str, ...]) -> int:
    for alias in aliases:
        key = _norm_header(alias)
        if key in norm_headers:
            return norm_headers.index(key)
    return -1


def extract_claim_rows(
    xlsx_bytes: bytes,
    *,
    claim_id_column: str | None = None,
    sheet_name: str | None = None,
) -> tuple[list[dict[str, Any]], str]:
    """Return ``(claim_rows, resolved_claim_id_column)``.

    Each row dict contains at minimum ``claim_id``. When present in the
    workbook, also extracts ``paid_dt``, ``total_billed``, and ``total_paid``
    (header matching is case/whitespace tolerant).
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise XlsxParseError("openpyxl is required to parse xlsx files") from exc

    try:
        wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise XlsxParseError(f"failed to open workbook: {exc}") from exc

    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    rows: Iterable = ws.iter_rows(values_only=True)
    header_row = next(rows, None)
    if not header_row:
        raise XlsxParseError("workbook has no rows")

    headers = [(str(h).strip() if h is not None else "") for h in header_row]
    norm = [_norm_header(h) for h in headers]

    wanted = _norm_header(claim_id_column) if claim_id_column else ""
    claim_idx = norm.index(wanted) if wanted and wanted in norm else -1
    if claim_idx < 0:
        for alias in _CLAIM_ID_ALIASES:
            claim_idx = _find_column(norm, (alias,))
            if claim_idx >= 0:
                break
    if claim_idx < 0:
        raise XlsxParseError(
            f"could not find a claim-id column (looked for {claim_id_column!r} "
            f"and aliases {list(_CLAIM_ID_ALIASES)}); headers found: {headers}"
        )
    resolved = headers[claim_idx]

    extra_idxs: dict[str, int] = {}
    for field, aliases in _OPTIONAL_COLUMNS.items():
        idx = _find_column(norm, aliases)
        if idx >= 0:
            extra_idxs[field] = idx

    auditor_idx = _find_column(norm, _AUDITOR_NAME_ALIASES)

    claim_rows: list[dict[str, Any]] = []
    for row in rows:
        if claim_idx >= len(row):
            continue
        claim_cell = row[claim_idx]
        if claim_cell is None:
            continue
        claim_id = str(claim_cell).strip()
        if not claim_id:
            continue

        record: dict[str, Any] = {"claim_id": claim_id}
        for field, idx in extra_idxs.items():
            if idx < len(row):
                value = _cell_value(row[idx])
                if value is not None:
                    record[field] = value
        if auditor_idx >= 0 and auditor_idx < len(row):
            auditor_name = _cell_value(row[auditor_idx])
            if auditor_name is not None:
                record["original_auditor"] = str(auditor_name).strip()
        claim_rows.append(record)

    return claim_rows, resolved


def extract_claim_ids(
    xlsx_bytes: bytes,
    *,
    claim_id_column: str | None = None,
    sheet_name: str | None = None,
) -> tuple[list[str], str]:
    """Return ``(claim_ids, resolved_column_name)`` — backward-compatible helper."""
    rows, resolved = extract_claim_rows(
        xlsx_bytes,
        claim_id_column=claim_id_column,
        sheet_name=sheet_name,
    )
    return [str(r["claim_id"]) for r in rows], resolved
